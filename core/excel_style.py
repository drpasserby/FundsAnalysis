# -*- coding: utf-8 -*-
"""xlsx 统一样式：表头加粗填色 + 冻结首行 + 按内容估算列宽。

性能约束：列宽估算用 `iter_rows` 且只扫前 max_scan_rows 行，
避免 `ws[letter]` 把整列物化（WehatPivotTable v1.0.0 的老写法）。
"""
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
HEADER_FONT = Font(bold=True)


def _display_width(value):
    """粗略显示宽度：中文按 2 计，其余按 1 计。"""
    w = 0
    for ch in str(value):
        w += 2 if ord(ch) > 127 else 1
    return w


def style_worksheet(ws, max_scan_rows=300, min_width=10, max_width=45):
    """统一表头样式 + 冻结首行 + 估算列宽。"""
    max_row = ws.max_row or 1
    max_col = ws.max_column or 1
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    scan_to = min(max_scan_rows, max_row)
    for i in range(1, max_col + 1):
        widest = 0
        for row in ws.iter_rows(min_row=1, max_row=scan_to, min_col=i, max_col=i):
            for cell in row:
                if cell.value is None:
                    continue
                w = _display_width(cell.value)
                if w > widest:
                    widest = w
        ws.column_dimensions[get_column_letter(i)].width = min(max(widest + 3, min_width), max_width)
