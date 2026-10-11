"""
微信流水生成数据透视表
版本：1.0.0
作者：wulvxinchen
"""
import os
import sys
import time
import traceback
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

# ================== 配置区 ==================
TARGET_FILENAME   = 'TenpayTrades.txt'
MERGED_TXT_NAME   = '总流水.txt'
SUMMARY_XLSX_NAME = '流水总表.xlsx'
PIVOT_XLSX_NAME   = '数据透视表.xlsx'

SHEET_RAW      = '流水总表'
SHEET_PIVOT    = '数据透视表'
SHEET_RELATION = '人物关系图'

COL_ACCOUNT      = '用户侧账号名称'
COL_LOAN_TYPE    = '借贷类型'
COL_COUNTERPARTY = '对手侧账户名称'
COL_AMOUNT       = '交易金额(分)'
PIVOT_AMOUNT     = '求和项:交易金额(分)'

COL_USER_ROLE = '用户侧角色'
COL_CP_ROLE   = '对手侧角色'

HEADER_FILL = PatternFill('solid', fgColor='DDEBF7')
HEADER_FONT = Font(bold=True)
# ============================================


def log(*args):
    """带时间戳的日志，直接输出到终端"""
    print('[%s]' % time.strftime('%H:%M:%S'), *args, flush=True)


# ---------------------------------------------------------------
# 0. 进度条窗口
# ---------------------------------------------------------------
class ProgressWindow:
    def __init__(self, parent, title='处理中...'):
        self.parent = parent
        self.win = tk.Toplevel(parent)
        self.win.title(title)
        self.win.resizable(False, False)
        self.win.protocol('WM_DELETE_WINDOW', lambda: None)

        W, H = 480, 150
        self.win.update_idletasks()
        sw = self.win.winfo_screenwidth()
        sh = self.win.winfo_screenheight()
        x = (sw - W) // 2
        y = (sh - H) // 2
        self.win.geometry('%dx%d+%d+%d' % (W, H, x, y))

        tk.Label(self.win, text=title,
                 font=('Microsoft YaHei', 11, 'bold')).pack(pady=(16, 6))

        self.status_var = tk.StringVar(value='准备中...')
        tk.Label(self.win, textvariable=self.status_var,
                 font=('Microsoft YaHei', 9), anchor='w',
                 justify='left', wraplength=440).pack(fill='x', padx=20)

        self.pb = ttk.Progressbar(self.win, orient='horizontal',
                                  length=440, mode='determinate', maximum=100)
        self.pb.pack(padx=20, pady=(8, 16))

        self._indeterminate = False
        self.win.lift()
        self.win.update()

    def set_indeterminate(self, text=None):
        self._indeterminate = True
        self.pb.configure(mode='indeterminate', maximum=100, value=0)
        self.pb.start(15)
        if text:
            self.status_var.set(text)
        self.win.update()

    def set_determinate(self, total, text=None):
        self._indeterminate = False
        try:
            self.pb.stop()
        except Exception:
            pass
        self.pb.configure(mode='determinate', maximum=max(total, 1), value=0)
        if text:
            self.status_var.set(text)
        self.win.update()

    def update(self, value, total=None, text=None):
        if self._indeterminate:
            try:
                self.pb.stop()
            except Exception:
                pass
            self._indeterminate = False
            self.pb.configure(mode='determinate')
        if total is not None:
            self.pb.configure(maximum=max(total, 1))
        self.pb.configure(value=value)
        if text:
            self.status_var.set(text)
        self.win.update()

    def set_status(self, text):
        self.status_var.set(text)
        self.win.update()

    def close(self):
        try:
            self.pb.stop()
        except Exception:
            pass
        try:
            self.win.destroy()
        except Exception:
            pass
        # 让主窗口处理一次残留事件，避免后续 Toplevel 无法显示
        try:
            self.parent.update()
        except Exception:
            pass


# ---------------------------------------------------------------
# 1. 选择文件夹
# ---------------------------------------------------------------
def choose_folders(root):
    folders = []
    while True:
        d = filedialog.askdirectory(
            parent=root,
            title='请选择包含 TenpayTrades.txt 的文件夹（点"取消"结束选择）'
        )
        if not d:
            break
        d = os.path.normpath(d)
        if d not in folders:
            folders.append(d)
        if not messagebox.askyesno(
                '继续选择',
                '已选择 %d 个文件夹：\n%s\n\n是否继续添加其他文件夹？'
                % (len(folders), '\n'.join(folders))):
            break
    return folders


# ---------------------------------------------------------------
# 2. 递归查找
# ---------------------------------------------------------------
def find_target_files(folders, progress=None):
    files = []
    for fi, folder in enumerate(folders, 1):
        msg = '正在扫描文件夹 (%d/%d)：%s' % (fi, len(folders), folder)
        log(msg)
        if progress:
            progress.set_status(msg)
        for dirpath, _dirnames, filenames in os.walk(folder):
            for fn in filenames:
                if fn == TARGET_FILENAME:
                    files.append(os.path.join(dirpath, fn))
    return sorted(set(files))


# ---------------------------------------------------------------
# 3. 读取单个文件
# ---------------------------------------------------------------
def read_one_file(path):
    try:
        with open(path, 'r', encoding='utf-8-sig', errors='replace') as f:
            raw = [ln.rstrip('\r\n') for ln in f]
    except Exception:
        try:
            with open(path, 'r', encoding='gbk', errors='replace') as f:
                raw = [ln.rstrip('\r\n') for ln in f]
        except Exception:
            return None, []

    header_idx = None
    for i, ln in enumerate(raw[:100]):
        if COL_ACCOUNT in ln and '\t' in ln:
            header_idx = i
            break
    if header_idx is None:
        for i, ln in enumerate(raw[:100]):
            if ln.count('\t') >= 2:
                header_idx = i
                break
    if header_idx is None:
        return None, []

    header = [h.strip() for h in raw[header_idx].split('\t')]
    n = len(header)

    rows = []
    for ln in raw[header_idx + 1:]:
        s = ln.strip()
        if not s:
            continue
        if s.startswith('---') or s.startswith('注') or s.startswith('说明'):
            continue
        parts = ln.split('\t')
        if len(parts) < 2:
            continue
        if len(parts) < n:
            parts = parts + [''] * (n - len(parts))
        elif len(parts) > n:
            parts = parts[:n]
        rows.append(parts)

    return header, rows


# ---------------------------------------------------------------
# 4. 列选择对话框（修复版：不依赖隐藏 root，强制置顶）
# ---------------------------------------------------------------
def select_columns(root, columns):
    log('打开列选择对话框，共 %d 列' % len(columns))
    result = {'ok': False, 'cols': []}

    # 关键修复：不要用 dlg.transient(root)，因为 root 已经 withdraw，
    # 某些平台上会导致对话框被"拖"成隐藏状态，从而 wait_window() 永远阻塞。
    dlg = tk.Toplevel()
    dlg.title('选择需要导入的列')

    W, H = 440, 560
    dlg.update_idletasks()
    sw = dlg.winfo_screenwidth()
    sh = dlg.winfo_screenheight()
    x = (sw - W) // 2
    y = (sh - H) // 2
    dlg.geometry('%dx%d+%d+%d' % (W, H, x, y))

    # 强制显示到最前
    dlg.deiconify()
    dlg.lift()
    dlg.attributes('-topmost', True)
    dlg.after(300, lambda: dlg.attributes('-topmost', False))
    dlg.focus_force()
    dlg.grab_set()

    tk.Label(dlg, text='请勾选需要导入 Excel 的列（默认全选）：',
             font=('Microsoft YaHei', 10)).pack(pady=8)

    container = tk.Frame(dlg)
    container.pack(fill='both', expand=True, padx=10)

    canvas = tk.Canvas(container, highlightthickness=0)
    sb = tk.Scrollbar(container, orient='vertical', command=canvas.yview)
    inner = tk.Frame(canvas)

    inner.bind('<Configure>',
               lambda e: canvas.configure(scrollregion=canvas.bbox('all')))
    canvas.create_window((0, 0), window=inner, anchor='nw')
    canvas.configure(yscrollcommand=sb.set)
    canvas.pack(side='left', fill='both', expand=True)
    sb.pack(side='right', fill='y')

    vars_ = []
    for c in columns:
        v = tk.BooleanVar(value=True)
        tk.Checkbutton(inner, text=c, variable=v, anchor='w').pack(fill='x')
        vars_.append((c, v))

    def on_all():
        for _, v in vars_:
            v.set(True)

    def on_none():
        for _, v in vars_:
            v.set(False)

    def on_ok():
        result['ok'] = True
        result['cols'] = [c for c, v in vars_ if v.get()]
        dlg.destroy()

    def on_cancel():
        result['ok'] = False
        dlg.destroy()

    btns = tk.Frame(dlg)
    btns.pack(fill='x', pady=8)
    tk.Button(btns, text='全选', width=8, command=on_all).pack(side='left', padx=6)
    tk.Button(btns, text='全不选', width=8, command=on_none).pack(side='left', padx=6)
    tk.Button(btns, text='确定', width=10, command=on_ok).pack(side='right', padx=6)
    tk.Button(btns, text='取消', width=10, command=on_cancel).pack(side='right', padx=6)

    dlg.wait_window()
    log('列选择对话框已关闭，选中 %d 列' % len(result['cols']))
    return result['cols'] if result['ok'] else None


# ---------------------------------------------------------------
# 5. 美化
# ---------------------------------------------------------------
def beautify_workbook(path, max_scan_rows=300):
    try:
        wb = openpyxl.load_workbook(path)
    except Exception as e:
        log('美化跳过（无法打开 %s）：%s' % (path, e))
        return
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = HEADER_FONT
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.freeze_panes = 'A2'
        for i in range(1, ws.max_column + 1):
            letter = get_column_letter(i)
            max_w = 0
            for r, cell in enumerate(ws[letter], start=1):
                if r > max_scan_rows:
                    break
                if cell.value is None:
                    continue
                txt = str(cell.value)
                w = sum(2 if ord(ch) > 127 else 1 for ch in txt)
                max_w = max(max_w, w)
            ws.column_dimensions[letter].width = min(max(max_w + 3, 10), 45)
    wb.save(path)


# ---------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------
def run():
    log('启动程序')
    root = tk.Tk()
    root.withdraw()

    # ---------- 选择文件夹 ----------
    log('等待用户选择文件夹...')
    folders = choose_folders(root)
    if not folders:
        log('用户未选择任何文件夹，退出')
        messagebox.showwarning('提示', '未选择任何文件夹，程序退出。', parent=root)
        return
    log('已选择 %d 个文件夹' % len(folders))

    # ==========================================================
    # 阶段 1：扫描 + 读取 + 合并
    # ==========================================================
    log('创建进度窗口')
    prog = ProgressWindow(root, '正在处理流水文件...')
    prog.set_indeterminate('正在扫描文件夹，查找 %s ...' % TARGET_FILENAME)

    files = find_target_files(folders, progress=prog)
    log('扫描完成，共找到 %d 个 %s' % (len(files), TARGET_FILENAME))

    if not files:
        prog.close()
        log('未找到目标文件，退出')
        messagebox.showwarning('提示',
                               '在所选文件夹中未找到任何 "%s" 文件。' % TARGET_FILENAME,
                               parent=root)
        return

    total_files = len(files)
    prog.set_determinate(total_files,
                         '共找到 %d 个文件，开始读取...' % total_files)

    base_header = None
    all_rows = []
    read_count = 0

    for i, p in enumerate(files, 1):
        short = p if len(p) <= 60 else '...' + p[-57:]
        prog.update(i - 1, total_files,
                    '读取文件 (%d/%d)：%s' % (i, total_files, short))
        header, rows = read_one_file(p)
        if header is not None:
            if base_header is None:
                base_header = header
            name2idx = {}
            for k, name in enumerate(header):
                if name not in name2idx:
                    name2idx[name] = k
            idx_list = [name2idx.get(c, None) for c in base_header]
            for r in rows:
                all_rows.append([r[k] if k is not None and k < len(r) else ''
                                 for k in idx_list])
            read_count += 1
        prog.update(i, total_files,
                    '读取文件 (%d/%d)：%s' % (i, total_files, short))

    log('读取完成：成功解析 %d 个文件，共 %d 行数据' % (read_count, len(all_rows)))

    if base_header is None or not all_rows:
        prog.close()
        log('解析不到有效数据，退出')
        messagebox.showerror('错误', '未能从这些文件中解析出有效数据。', parent=root)
        return

    out_dir = folders[0]

    # ---------- 写出 总流水.txt ----------
    merged_path = os.path.join(out_dir, MERGED_TXT_NAME)
    total_lines = len(all_rows)
    prog.set_determinate(total_lines, '正在写出 %s ...' % MERGED_TXT_NAME)
    log('开始写出 %s（共 %d 行）' % (merged_path, total_lines))

    chunk = max(1, total_lines // 200)
    with open(merged_path, 'w', encoding='utf-8-sig', newline='') as f:
        f.write('\t'.join(base_header) + '\n')
        for i, r in enumerate(all_rows, 1):
            f.write('\t'.join(r) + '\n')
            if i % chunk == 0 or i == total_lines:
                prog.update(i, total_lines,
                            '写出 %s (%d/%d 行)'
                            % (MERGED_TXT_NAME, i, total_lines))

    log('总流水.txt 写出完毕：%s' % merged_path)
    prog.close()
    log('进度窗口已关闭')

    # 让主线程处理一次事件，避免后续 Toplevel 显示异常
    try:
        root.update()
    except Exception:
        pass

    # ==========================================================
    # 阶段 2：选择需要导入的列
    # ==========================================================
    log('构建 DataFrame ...')
    df = pd.DataFrame(all_rows, columns=base_header)
    log('DataFrame 构建完成，shape = %s' % (df.shape,))

    # 用非正则方式快速去空格（大文件下 regex 版本极慢）
    log('清理空白字符 ...')
    for col in df.columns:
        try:
            df[col] = df[col].astype(str).str.strip()
        except Exception:
            pass
    log('清理完成')

    selected = select_columns(root, list(df.columns))
    if not selected:
        log('用户取消列选择，退出')
        messagebox.showwarning('提示', '未选择任何列，程序退出。', parent=root)
        return
    df_sel = df[selected].copy()

    need_cols = [COL_ACCOUNT, COL_LOAN_TYPE, COL_COUNTERPARTY, COL_AMOUNT]
    missing = [c for c in need_cols if c not in df_sel.columns]
    if missing:
        log('缺少必要列：%s' % missing)
        messagebox.showerror(
            '缺少必要列',
            '生成数据透视表需要以下列，请重新运行并勾选：\n\n' + '\n'.join(missing),
            parent=root)
        return

    # ==========================================================
    # 阶段 3：生成 Excel
    # ==========================================================
    log('创建第二个进度窗口（Excel 阶段）')
    prog = ProgressWindow(root, '正在生成 Excel 文件...')
    prog.set_determinate(100, '正在准备数据...')

    log('转换"交易金额(分)"为数值 ...')
    prog.update(5, 100, '正在转换"交易金额(分)"为数值...')
    df_sel[COL_AMOUNT] = (
        df_sel[COL_AMOUNT].astype(str)
        .str.replace(',', '', regex=False)
        .str.replace('"', '', regex=False)
        .str.strip()
    )
    df_sel[COL_AMOUNT] = pd.to_numeric(df_sel[COL_AMOUNT], errors='coerce').fillna(0)
    log('金额转换完成')

    log('生成数据透视表内容 ...')
    prog.update(15, 100, '正在生成数据透视表内容...')
    try:
        pivot = (
            df_sel.groupby([COL_ACCOUNT, COL_LOAN_TYPE, COL_COUNTERPARTY],
                           dropna=False, as_index=False, sort=True)[COL_AMOUNT]
            .sum()
        )
    except Exception:
        pivot = (
            df_sel.groupby([COL_ACCOUNT, COL_LOAN_TYPE, COL_COUNTERPARTY],
                           dropna=False, as_index=False, sort=False)[COL_AMOUNT]
            .sum()
        )
    pivot = pivot.rename(columns={COL_AMOUNT: PIVOT_AMOUNT})
    pivot = pivot[[COL_ACCOUNT, COL_LOAN_TYPE, COL_COUNTERPARTY, PIVOT_AMOUNT]]
    log('透视表内容生成完成，共 %d 行' % len(pivot))

    # 写出 流水总表.xlsx
    summary_path = os.path.join(out_dir, SUMMARY_XLSX_NAME)
    log('写出 %s ...' % summary_path)
    prog.update(30, 100, '正在写出 %s ...' % SUMMARY_XLSX_NAME)
    with pd.ExcelWriter(summary_path, engine='openpyxl') as writer:
        df_sel.to_excel(writer, sheet_name=SHEET_RAW, index=False)
        prog.update(55, 100, '正在写入工作表"%s"...' % SHEET_PIVOT)
        pivot.to_excel(writer, sheet_name=SHEET_PIVOT, index=False)
    log('流水总表.xlsx 写出完毕')

    # 写出 数据透视表.xlsx
    pivot_path = os.path.join(out_dir, PIVOT_XLSX_NAME)
    log('写出 %s ...' % pivot_path)
    prog.update(65, 100, '正在写出 %s ...' % PIVOT_XLSX_NAME)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_PIVOT
    ws.append([COL_ACCOUNT, COL_USER_ROLE, COL_LOAN_TYPE,
               COL_COUNTERPARTY, COL_CP_ROLE, PIVOT_AMOUNT])

    data_rows = pivot.values.tolist()
    total_pr = len(data_rows)
    chunk_pr = max(1, total_pr // 100)
    for r, row in enumerate(data_rows, start=2):
        ws.cell(row=r, column=1, value=row[0])
        ws.cell(row=r, column=2,
                value='=VLOOKUP(A%d,%s!B:C,2,0)' % (r, SHEET_RELATION))
        ws.cell(row=r, column=3, value=row[1])
        ws.cell(row=r, column=4, value=row[2])
        ws.cell(row=r, column=5,
                value='=VLOOKUP(D%d,%s!B:C,2,0)' % (r, SHEET_RELATION))
        ws.cell(row=r, column=6, value=row[3])
        if (r - 1) % chunk_pr == 0 or (r - 1) == total_pr:
            prog.update(65 + int(15 * (r - 1) / max(total_pr, 1)), 100,
                        '写入透视表及公式 (%d/%d)' % (r - 1, total_pr))

    log('创建"%s"工作表 ...' % SHEET_RELATION)
    prog.update(82, 100, '正在创建工作表"%s"...' % SHEET_RELATION)
    ws_rel = wb.create_sheet(SHEET_RELATION)
    ws_rel.append(['序号', '姓名', '角色'])
    wb.save(pivot_path)
    log('数据透视表.xlsx 写出完毕')

    log('美化 %s ...' % SUMMARY_XLSX_NAME)
    prog.update(88, 100, '正在美化 %s ...' % SUMMARY_XLSX_NAME)
    beautify_workbook(summary_path)
    log('美化 %s ...' % PIVOT_XLSX_NAME)
    prog.update(95, 100, '正在美化 %s ...' % PIVOT_XLSX_NAME)
    beautify_workbook(pivot_path)

    prog.update(100, 100, '处理完成！')
    prog.close()
    log('全部处理完成')

    msg = (
        '处理完成！\n\n'
        '共扫描到 %d 个 %s 文件，成功解析 %d 个。\n'
        '数据总行数：%d\n\n'
        '输出文件目录：\n%s\n\n'
        '1) %s\n'
        '2) %s（工作表：%s、%s）\n'
        '3) %s（工作表：%s、%s）\n\n'
        '提示：打开 %s 后，请在"%s"工作表中填写"序号/姓名/角色"，\n'
        'VLOOKUP 公式即可自动带出角色。'
        % (len(files), TARGET_FILENAME, read_count, len(all_rows),
           out_dir,
           MERGED_TXT_NAME,
           SUMMARY_XLSX_NAME, SHEET_RAW, SHEET_PIVOT,
           PIVOT_XLSX_NAME, SHEET_PIVOT, SHEET_RELATION,
           PIVOT_XLSX_NAME, SHEET_RELATION)
    )
    messagebox.showinfo('完成', msg, parent=root)
    root.destroy()
    log('退出')


if __name__ == '__main__':
    try:
        run()
    except Exception:
        err = traceback.format_exc()
        print(err, flush=True)
        try:
            r = tk.Tk()
            r.withdraw()
            messagebox.showerror('程序异常', err)
            r.destroy()
        except Exception:
            pass